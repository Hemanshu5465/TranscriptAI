"""Transcript exporters: txt, docx, pdf, srt, vtt, json, csv."""
from __future__ import annotations

import csv
import io
import json
from dataclasses import dataclass, field

from app.utils.timecode import hhmmss_full, srt_timestamp, vtt_timestamp

SUPPORTED_FORMATS = ("txt", "docx", "pdf", "srt", "vtt", "json", "csv")

_MEDIA_TYPES = {
    "txt": "text/plain; charset=utf-8",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pdf": "application/pdf",
    "srt": "application/x-subrip; charset=utf-8",
    "vtt": "text/vtt; charset=utf-8",
    "json": "application/json; charset=utf-8",
    "csv": "text/csv; charset=utf-8",
}


@dataclass
class ExportWord:
    word: str
    start: float
    end: float
    confidence: float | None = None


@dataclass
class ExportSegment:
    index: int
    start: float
    end: float
    text: str
    speaker: str | None = None
    confidence: float | None = None
    words: list[ExportWord] = field(default_factory=list)


@dataclass
class ExportBundle:
    title: str
    video_id: str
    url: str
    channel: str | None
    language: str | None
    duration_seconds: float | None
    segments: list[ExportSegment]
    full_text: str
    stats: dict = field(default_factory=dict)


@dataclass
class ExportFile:
    content: bytes
    media_type: str
    filename: str


def _safe_slug(text: str) -> str:
    keep = "".join(c if c.isalnum() or c in " -_" else "" for c in text).strip()
    return (keep.replace(" ", "-")[:60] or "transcript").lower()


def export(bundle: ExportBundle, fmt: str, *, include_timestamps: bool = True) -> ExportFile:
    fmt = fmt.lower()
    if fmt not in SUPPORTED_FORMATS:
        raise ValueError(f"Unsupported export format: {fmt!r}")
    content = _DISPATCH[fmt](bundle, include_timestamps)
    return ExportFile(
        content=content,
        media_type=_MEDIA_TYPES[fmt],
        filename=f"{_safe_slug(bundle.title)}.{fmt}",
    )


# --- individual writers -------------------------------------------------------

def _as_txt(b: ExportBundle, include_timestamps: bool) -> bytes:
    lines: list[str] = [b.title, b.url, ""]
    if include_timestamps:
        for seg in b.segments:
            prefix = f"[{hhmmss_full(seg.start)}]"
            if seg.speaker:
                prefix += f" {seg.speaker}:"
            lines.append(prefix)
            lines.append(seg.text)
            lines.append("")
    else:
        lines.append(b.full_text)
    return ("\n".join(lines).rstrip() + "\n").encode("utf-8")


def _as_srt(b: ExportBundle, include_timestamps: bool) -> bytes:
    out: list[str] = []
    for n, seg in enumerate(b.segments, start=1):
        text = f"{seg.speaker}: {seg.text}" if seg.speaker else seg.text
        out.append(str(n))
        out.append(f"{srt_timestamp(seg.start)} --> {srt_timestamp(seg.end)}")
        out.append(text)
        out.append("")
    return ("\n".join(out)).encode("utf-8")


def _as_vtt(b: ExportBundle, include_timestamps: bool) -> bytes:
    out: list[str] = ["WEBVTT", ""]
    for n, seg in enumerate(b.segments, start=1):
        text = f"<v {seg.speaker}>{seg.text}" if seg.speaker else seg.text
        out.append(str(n))
        out.append(f"{vtt_timestamp(seg.start)} --> {vtt_timestamp(seg.end)}")
        out.append(text)
        out.append("")
    return ("\n".join(out)).encode("utf-8")


def _as_json(b: ExportBundle, include_timestamps: bool) -> bytes:
    payload = {
        "video": {
            "id": b.video_id,
            "url": b.url,
            "title": b.title,
            "channel": b.channel,
        },
        "duration": b.duration_seconds,
        "language": b.language,
        "stats": b.stats,
        "full_text": b.full_text,
        "segments": [
            {
                "index": s.index,
                "start": s.start,
                "end": s.end,
                "speaker": s.speaker,
                "confidence": s.confidence,
                "text": s.text,
            }
            for s in b.segments
        ],
        "words": [
            {
                "segment_index": s.index,
                "word": w.word,
                "start": w.start,
                "end": w.end,
                "confidence": w.confidence,
            }
            for s in b.segments
            for w in s.words
        ],
    }
    return json.dumps(payload, ensure_ascii=False, indent=2).encode("utf-8")


def _as_csv(b: ExportBundle, include_timestamps: bool) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(["index", "start", "end", "speaker", "confidence", "text"])
    for s in b.segments:
        writer.writerow(
            [
                s.index,
                f"{s.start:.3f}",
                f"{s.end:.3f}",
                s.speaker or "",
                "" if s.confidence is None else f"{s.confidence:.3f}",
                s.text,
            ]
        )
    return buf.getvalue().encode("utf-8")


def _as_docx(b: ExportBundle, include_timestamps: bool) -> bytes:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.add_heading(b.title or "Transcript", level=1)
    meta = doc.add_paragraph()
    meta.add_run(f"{b.channel or 'Unknown channel'}  •  {b.url}").italic = True
    if b.language:
        doc.add_paragraph(f"Language: {b.language}").runs[0].italic = True
    doc.add_paragraph("")

    if include_timestamps:
        for seg in b.segments:
            head = doc.add_paragraph()
            run = head.add_run(
                f"[{hhmmss_full(seg.start)}]"
                + (f"  {seg.speaker}" if seg.speaker else "")
            )
            run.bold = True
            run.font.size = Pt(10)
            doc.add_paragraph(seg.text)
    else:
        for para in b.full_text.split("\n\n"):
            doc.add_paragraph(para)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _as_pdf(b: ExportBundle, include_timestamps: bool) -> bytes:
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=LETTER,
        leftMargin=0.9 * inch, rightMargin=0.9 * inch,
        topMargin=0.9 * inch, bottomMargin=0.9 * inch,
        title=b.title or "Transcript",
    )
    styles = getSampleStyleSheet()
    ts_style = ParagraphStyle(
        "ts", parent=styles["Normal"], fontSize=8, textColor="#6366f1", spaceBefore=10
    )
    body = ParagraphStyle("body", parent=styles["Normal"], fontSize=10.5, leading=15)

    def esc(t: str) -> str:
        return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    story = [Paragraph(esc(b.title or "Transcript"), styles["Title"])]
    story.append(Paragraph(esc(f"{b.channel or ''}  •  {b.url}"), styles["Italic"]))
    story.append(Spacer(1, 14))

    if include_timestamps:
        for seg in b.segments:
            label = f"[{hhmmss_full(seg.start)}]" + (f"  {seg.speaker}" if seg.speaker else "")
            story.append(Paragraph(esc(label), ts_style))
            story.append(Paragraph(esc(seg.text), body))
    else:
        for para in b.full_text.split("\n\n"):
            story.append(Paragraph(esc(para), body))
            story.append(Spacer(1, 8))

    doc.build(story)
    return buf.getvalue()


_DISPATCH = {
    "txt": _as_txt,
    "docx": _as_docx,
    "pdf": _as_pdf,
    "srt": _as_srt,
    "vtt": _as_vtt,
    "json": _as_json,
    "csv": _as_csv,
}
